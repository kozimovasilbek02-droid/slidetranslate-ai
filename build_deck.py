# -*- coding: utf-8 -*-
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    PRIMARY = RGBColor(11, 61, 145)       # Deep Uzbek Blue
    SECONDARY = RGBColor(0, 150, 136)     # Emerald / Teal
    ACCENT = RGBColor(212, 175, 55)       # Gold
    DARK_TEXT = RGBColor(33, 37, 41)      # Dark Slate
    MUTED_TEXT = RGBColor(90, 100, 110)   # Muted Gray
    LIGHT_BG = RGBColor(245, 247, 250)    # Soft Gray-Blue Background
    WHITE = RGBColor(255, 255, 255)
    CARD_BG = RGBColor(255, 255, 255)
    BORDER_COLOR = RGBColor(220, 225, 232)

    TOTAL_SLIDES = 16

    def set_slide_bg(slide, color=LIGHT_BG):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = color
        bg.line.fill.background()
        slide.shapes._spTree.remove(bg._element)
        slide.shapes._spTree.insert(2, bg._element)

    def add_header(slide, title_text, category_text="O'ZBEKISTONNING ENG YANGI TARIXI"):
        tb = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.733), Inches(1.1))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        
        p_cat = tf.paragraphs[0]
        p_cat.text = category_text.upper()
        p_cat.font.size = Pt(11)
        p_cat.font.bold = True
        p_cat.font.color.rgb = SECONDARY
        p_cat.font.name = 'Arial'

        p_title = tf.add_paragraph()
        p_title.text = title_text
        p_title.font.size = Pt(23)
        p_title.font.bold = True
        p_title.font.color.rgb = PRIMARY
        p_title.font.name = 'Arial'
        p_title.space_before = Pt(3)

    def add_footer(slide, current_slide):
        line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(6.8), Inches(11.733), Inches(0.015))
        line.fill.solid()
        line.fill.fore_color.rgb = BORDER_COLOR
        line.line.fill.background()

        tb = slide.shapes.add_textbox(Inches(0.8), Inches(6.9), Inches(11.733), Inches(0.4))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.text = f"Yangi O'zbekiston Taraqqiyot Strategiyasi va Inson Qadri Tamoyili | {current_slide} / {TOTAL_SLIDES}"
        p.font.size = Pt(9.5)
        p.font.color.rgb = MUTED_TEXT
        p.font.name = 'Arial'

    def create_cards_layout(slide, title, category, cards_data, columns=3):
        set_slide_bg(slide)
        add_header(slide, title, category)
        
        num_cards = len(cards_data)
        rows = (num_cards + columns - 1) // columns
        
        start_x = Inches(0.8)
        start_y = Inches(1.6)
        total_width = Inches(11.733)
        total_height = Inches(4.9)
        
        gap_x = Inches(0.25)
        gap_y = Inches(0.25)
        
        card_width = (total_width - (gap_x * (columns - 1))) / columns
        card_height = (total_height - (gap_y * (rows - 1))) / rows
        
        for idx, card in enumerate(cards_data):
            r = idx // columns
            c = idx % columns
            
            x = start_x + c * (card_width + gap_x)
            y = start_y + r * (card_height + gap_y)
            
            card_shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, card_width, card_height)
            card_shape.fill.solid()
            card_shape.fill.fore_color.rgb = CARD_BG
            card_shape.line.color.rgb = BORDER_COLOR
            card_shape.line.width = Pt(1)
            
            accent_bar = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, card_width, Inches(0.08))
            accent_bar.fill.solid()
            accent_bar.fill.fore_color.rgb = card.get("accent_color", PRIMARY)
            accent_bar.line.fill.background()
            
            tb = slide.shapes.add_textbox(x + Inches(0.25), y + Inches(0.2), card_width - Inches(0.5), card_height - Inches(0.35))
            tf = tb.text_frame
            tf.word_wrap = True
            tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
            
            if card.get("tag"):
                p_tag = tf.paragraphs[0]
                p_tag.text = card["tag"].upper()
                p_tag.font.size = Pt(9)
                p_tag.font.bold = True
                p_tag.font.color.rgb = card.get("accent_color", SECONDARY)
                p_tag.font.name = 'Arial'
                p_head = tf.add_paragraph()
            else:
                p_head = tf.paragraphs[0]
                
            p_head.text = card["title"]
            p_head.font.size = Pt(14)
            p_head.font.bold = True
            p_head.font.color.rgb = PRIMARY
            p_head.font.name = 'Arial'
            p_head.space_before = Pt(2)
            
            for item in card.get("items", []):
                p_item = tf.add_paragraph()
                p_item.text = f"• {item}"
                p_item.font.size = Pt(11)
                p_item.font.color.rgb = DARK_TEXT
                p_item.font.name = 'Arial'
                p_item.space_before = Pt(5)

    # SLIDE 1: COVER
    slide1 = prs.slides.add_slide(blank_layout)
    bg1 = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg1.fill.solid()
    bg1.fill.fore_color.rgb = PRIMARY
    bg1.line.fill.background()

    acc1 = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.3), Inches(0.15), Inches(4.8))
    acc1.fill.solid()
    acc1.fill.fore_color.rgb = ACCENT
    acc1.line.fill.background()

    tb1 = slide1.shapes.add_textbox(Inches(1.2), Inches(1.3), Inches(11.0), Inches(4.8))
    tf1 = tb1.text_frame
    tf1.word_wrap = True

    p = tf1.paragraphs[0]
    p.text = "O'ZBEKISTONNING ENG YANGI TARIXI"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = ACCENT
    p.font.name = 'Arial'

    p2 = tf1.add_paragraph()
    p2.text = "Yangi O'zbekiston Taraqqiyot Strategiyasi va Inson Qadri Tamoyili"
    p2.font.size = Pt(32)
    p2.font.bold = True
    p2.font.color.rgb = WHITE
    p2.font.name = 'Arial'
    p2.space_before = Pt(12)

    p3 = tf1.add_paragraph()
    p3.text = "Inson qadrini ulug'lash, erkin fuqarolik jamiyatini barpo etish va milliy iqtisodiyotni yuksaltirishga qaratilgan strategik islohotlar tahlili"
    p3.font.size = Pt(16)
    p3.font.color.rgb = RGBColor(205, 225, 250)
    p3.font.name = 'Arial'
    p3.space_before = Pt(18)

    p4 = tf1.add_paragraph()
    p4.text = "Fan: Tarix / Ijtimoiy fanlar  |  Hajmi: 16 Slayd  |  Mustaqil ish"
    p4.font.size = Pt(12.5)
    p4.font.color.rgb = WHITE
    p4.font.name = 'Arial'
    p4.space_before = Pt(45)

    # SLIDE 2: MUNDARIJA
    slide2 = prs.slides.add_slide(blank_layout)
    set_slide_bg(slide2)
    add_header(slide2, "Taqdimot Rejasi (Mundarija)", "REJA VA TUZILISh")
    add_footer(slide2, 2)

    agenda_items = [
        {"num": "01", "title": "Yangi O'zbekiston tushunchasi va strategik maqsadlar", "desc": "Milliy tiklanishdan milliy yuksalishga o'tish bosqichlari"},
        {"num": "02", "title": "Inson qadri va uning davlat siyosatidagi ustuvor o'rni", "desc": "Inson - jamiyat - davlat tamoyili va ijtimoiy kafolatlar"},
        {"num": "03", "title": "Iqtisodiy islohotlar: erkin bozor va investitsion muhit", "desc": "Tadbirkorlikni qo'llab-quvvatlash va raqamli iqtisodiyot"},
        {"num": "04", "title": "Ta'lim va sog'liqni saqlash sohasidagi tub o'zgarishlar", "desc": "Inson kapitaliga investitsiya va tibbiy xizmat kalitlari"},
        {"num": "05", "title": "Tashqi siyosat va Markaziy Osiyodagi yangi muhit", "desc": "Ochoq, pragmatik va yaxshi qo'shnichilik diplomatiyasi"}
    ]

    for i, ag in enumerate(agenda_items):
        y_pos = Inches(1.6) + i * Inches(0.98)
        
        card = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), y_pos, Inches(11.733), Inches(0.85))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = BORDER_COLOR
        
        num_box = slide2.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.0), y_pos + Inches(0.12), Inches(0.6), Inches(0.6))
        num_box.fill.solid()
        num_box.fill.fore_color.rgb = PRIMARY
        num_box.line.fill.background()
        tf_num = num_box.text_frame
        p_num = tf_num.paragraphs[0]
        p_num.text = ag["num"]
        p_num.font.size = Pt(14)
        p_num.font.bold = True
        p_num.font.color.rgb = WHITE
        p_num.alignment = PP_ALIGN.CENTER
        
        tb = slide2.shapes.add_textbox(Inches(1.8), y_pos + Inches(0.1), Inches(10.5), Inches(0.65))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        
        p_t = tf.paragraphs[0]
        p_t.text = ag["title"]
        p_t.font.size = Pt(14)
        p_t.font.bold = True
        p_t.font.color.rgb = PRIMARY
        
        p_d = tf.add_paragraph()
        p_d.text = ag["desc"]
        p_d.font.size = Pt(11)
        p_d.font.color.rgb = MUTED_TEXT

    # SLIDE 3
    cards3 = [
        {"tag": "Tarixiy Kontekst", "title": "Tarixiy Burilish Bosqichi", "accent_color": PRIMARY, "items": ["2016-yildan e'tiboran O'zbekiston rivojlanishining mutlaqo yangi bosqichi boshlandi.", "Tarixiy taraqqiyotda davlat va jamiyat o'rtasidagi munosabatlar tubdan o'zgardi.", "Eski byurokratik g'ovlar olib tashlanib, islohotlarning ochiqligi ta'minlandi."]},
        {"tag": "Konsepsiya", "title": "Tarixiy Konsepsiya", "accent_color": SECONDARY, "items": ["'Milliy tiklanishdan - Milliy yuksalish sari' g'oyasi bosh tamoyilga aylandi.", "Islohotlar faqat raqamlar uchun emas, balki har bir fuqaro hayotida sezilishi shart qilindi.", "Tarixiy tajriba va zamonaviy xalqaro standartlar uyg'unlashtirildi."]},
        {"tag": "Maqsad", "title": "Bosh Maqsad va Vazifa", "accent_color": ACCENT, "items": ["Erkin, obod va farovon Yangi O'zbekistonni barpo etish.", "Dunyo hamjamiyatida mamlakatning munosib o'rni va nufuzini ta'minlash.", "Kelajak avlod uchun barqaror rivojlanish poydevorini yaratish."]}
    ]
    slide3 = prs.slides.add_slide(blank_layout)
    create_cards_layout(slide3, "Yangi Davrning Tarixiy Zaruriyati va Bosh G'oyasi", "KIRISH VA KONSEP-SIYA", cards3)
    add_footer(slide3, 3)

    # SLIDE 4
    cards4 = [
        {"title": "Strategik Yo'nalishlar", "accent_color": PRIMARY, "items": ["Erkin fuqarolik jamiyatini rivojlantirish.", "Adolat va qonun ustuvorligini ta'minlash.", "Milliy iqtisodiyotni jadal rivojlantirish.", "Adolatli ijtimoiy siyosat yuritish."]},
        {"title": "Asosiy Ustuvor Maqsadlar", "accent_color": SECONDARY, "items": ["Ma'naviy taraqqiyotni ta'minlash va sohani yangi bosqichga olib chiqish.", "Milliy manfaatlardan kelib chiqqan holda tashqi siyosat yuritish.", "Mamlakat xavfsizligi va mudofaa salohiyatini kuchaytirish."]},
        {"title": "Kutilayotgan Natijalar", "accent_color": ACCENT, "items": ["Aholi jon boshiga yalpi ichki mahsulot hajmini oshirish.", "O'rta daromadli mamlakatlar qatoridan munosib o'rin egallash.", "Kambag'allik darajasini kamida 2 baravarga qisqartirish."]}
    ]
    slide4 = prs.slides.add_slide(blank_layout)
    create_cards_layout(slide4, "1. Yangi O'zbekiston Taraqqiyot Strategiyasi (2022-2026)", "REJA 1: STRATEGIK MAQSADLAR", cards4)
    add_footer(slide4, 4)

    # SLIDE 5
    cards5 = [
        {"tag": "Transformatsiya", "title": "Boshqaruv Tizimi Islohoti", "accent_color": PRIMARY, "items": ["Ixcham va samarali davlat boshqaruv apparatini shakllantirish.", "Vazirlik va idoralar mas'uliyatini oshirish va maqbullashtirish.", "Davlat xizmatlarini to'liq raqamlashtirish va shaffoflikni ta'minlash."]},
        {"tag": "Mahallabay", "title": "Mahalla Tizimining O'rni", "accent_color": SECONDARY, "items": ["Mahalla - jamiyatning tayanch bo'g'iniga aylantirildi.", "'Hokim yordamchisi', 'Yoshlar yetakchisi' va 'Xotin-qizlar faoli' institutlari joriy etildi.", "Muammolarni bevosita joyida hal etish tizimi yo'lga qo'yildi."]},
        {"tag": "Raqamlashtirish", "title": "Xalq Bilan Muloqot", "accent_color": ACCENT, "items": ["Prezident Xalq qabulxonalari faoliyatining samaradorligi.", "Elektron hukumat (my.gov.uz) orqali yuzlab xizmatlarning yo'lga qo'yilishi.", "Jamoatchilik nazorati va fuqarolar tashabbusining oshishi."]}
    ]
    slide5 = prs.slides.add_slide(blank_layout)
    create_cards_layout(slide5, "Davlat Boshqaruvi va Mahalla Tizimida Transformatsiya", "REJA 1: DAVLAT BOSHQARUVI", cards5)
    add_footer(slide5, 5)

    # SLIDE 6
    cards6 = [
        {"tag": "Tamoyil 1", "title": "Inson Qadri - Bosh Mezon", "accent_color": PRIMARY, "items": ["'Inson - jamiyat - davlat' tamoyili ustuvor etib belgilandi.", "Davlat idoralari xalqqa xizmat qilishi shartligi konstitutsiyaviy darajada mustahkamlandi.", "Har bir fuqaro huquqi va erkinligi oliy qadriyat sanaladi."]},
        {"tag": "Tamoyil 2", "title": "Ijtimoiy Adolat va Himoya", "accent_color": SECONDARY, "items": ["Kam ta'minlangan va ehtiyojmand oilalarni manzilli qo'llab-quvvatlash.", "'Ayollar daftari', 'Yoshlar daftari' va 'Saxovat daftari' tizimlari samaradorligi.", "Nogironligi bo'lgan shaxslarga mehir va e'tibor qaratish."]},
        {"tag": "Tamoyil 3", "title": "Konstitutsiyaviy Islohotlar", "accent_color": ACCENT, "items": ["Yangi tahrirdagi Konstitutsiyada 'Ijtimoiy Davlat' maqomining belgilanishi.", "Insonning mehnati, yashashi va ta'lim olishiga bo'lgan huquqlarining kengaytirilishi.", "Shaxsiy daxlsizlik va xususiy mulk kafolatlarining kuchaytirilishi."]}
    ]
    slide6 = prs.slides.add_slide(blank_layout)
    create_cards_layout(slide6, "2. 'Inson Qadri Uchun' Tamoyilining Mazmun-Mohiyati", "REJA 2: INSON QADRI", cards6)
    add_footer(slide6, 6)

    # SLIDE 7
    cards7 = [
        {"title": "Sud-Huquq Tizimi Islohoti", "accent_color": PRIMARY, "items": ["Sudlarning haqiqiy mustaqilligini ta'minlash.", "Inson huquqlarini poymol etuvchi har qanday harakatga barham berish.", "Advokatura institutini va himoya huquqini kuchaytirish."]},
        {"title": "Majburiy Mehnatga Barham", "accent_color": SECONDARY, "items": ["Bolalar mehnati va majburiy mehnat to'liq tugatildi.", "Xalqaro Mehnat Tashkiloti (XMT) tomonidan e'tirof etildi.", "Paxta boykoti (Cotton Campaign) bekor qilinishiga erishildi."]},
        {"title": "Fikr va So'z Erkinligi", "accent_color": ACCENT, "items": ["Ommaviy axborot vositalari (OAV) va blogerlar faoliyatiga keng imkoniyatlar.", "Davlat organlari faoliyatining ochiqligi va shaffofligi.", "Jamoatchilik fikrining davlat qarorlariga ta'sir o'tkazishi."]}
    ]
    slide7 = prs.slides.add_slide(blank_layout)
    create_cards_layout(slide7, "Inson Huquqlari, Sud Mustaqilligi va So'z Erkinligi", "REJA 2: HUQUQIY KAFOLATLAR", cards7)
    add_footer(slide7, 7)

    # SLIDE 8
    cards8 = [
        {"tag": "Valyuta Islohoti", "title": "Valyuta Bozorini Erkinlashtirish", "accent_color": PRIMARY, "items": ["2017-yildagi erkin valyuta konvertatsiyasining joriy etilishi.", "Xorijiy investorlar va mahalliy tadbirkorlar uchun teng sharoitlar yaratish.", "Valyuta cheklovlarining bekor qilinishi."]},
        {"tag": "Soliq Tizimi", "title": "Soliq Yukini Kamaytirish", "accent_color": SECONDARY, "items": ["Soliq stavkalarining maqbullashtirilishi va soddalashtirilishi.", "QQS (Qo'shilgan qiymat solig'i) stavkasining 15% dan 12% ga tushirilishi.", "Halol tadbirkorlik uchun imtiyozli sharoitlar."]},
        {"tag": "Erkin Bozor", "title": "Mulkchilik va Xususiylashtirish", "accent_color": ACCENT, "items": ["Xususiy mulk daxlsizligining qonuniy kafolatlanishi.", "Davlat aktivlarini shaffof tenderlar orqali xususiylashtirish.", "Monopoliyani kamaytirish va raqobat muhitini rivojlantirish."]}
    ]
    slide8 = prs.slides.add_slide(blank_layout)
    create_cards_layout(slide8, "3. Iqtisodiy Islohotlar: Valyuta, Soliq va Erkin Bozor", "REJA 3: IQTISODIYOT", cards8)
    add_footer(slide8, 8)

    # SLIDE 9
    cards9 = [
        {"title": "Investitsiya Muhiti", "accent_color": PRIMARY, "items": ["Xorijiy sarmoyadorlar uchun huquqiy va iqtisodiy kafolatlar.", "Maxsus iqtisodiy zonalar (MIZ) va texnoparklar tarmog'ini kengaytirish.", "Xalqaro kredit reytinqlarida O'zbekiston pozitsiyasining yaxshilanishi."]},
        {"title": "Kichik va O'rta Biznes", "accent_color": SECONDARY, "items": ["Tadbirkorlik faoliyatini tekshirishlarga moratoriy va yengilliklar.", "Imtiyozli kreditlar va subsidiyalar ajratish hajmining oshishi.", "Yangi ish o'rinlarini yaratishda biznesning asosiy drayverga aylanishi."]},
        {"title": "Sanoat va Eksport", "accent_color": ACCENT, "items": ["Xomashyo eksportidan tayyor mahsulot eksportiga o'tish strategiyasi.", "Avtomobilsozlik, to'qimachilik va kimyo sanoatida modernizatsiya.", "Jahon Savdo Tashkilotiga (JST) a'zo bo'lish bo'yicha jadal muzokaralar."]}
    ]
    slide9 = prs.slides.add_slide(blank_layout)
    create_cards_layout(slide9, "Investitsiyalar Jalb Etish va Tadbirkorlikni Qo'llab-quvvatlash", "REJA 3: INVESTITSIYA VA BIZNES", cards9)
    add_footer(slide9, 9)

    # SLIDE 10
    cards10 = [
        {"tag": "Maktabgacha", "title": "Bog'cha Tizimi Islohoti", "accent_color": PRIMARY, "items": ["Maktabgacha ta'lim vazirligining tashkil etilishi.", "Qamrov darajasini 27% dan 72% dan ortiq ko'rsatkichga etkazish.", "Xususiy va davlat-xususiy sheriklik bog'chalarini ko'paytirish."]},
        {"tag": "Maktab Ta'limi", "title": "Maktablar va O'qituvchi Qadri", "accent_color": SECONDARY, "items": ["O'qituvchilar ish haqini oshirish va ularni majburiy mehnatsiz qilish.", "Prezident maktablari va Ijod maktablari tarmog'ini yaratish.", "Zamonaviy darsliklar va milliy o'quv dasturlarini joriy etish."]},
        {"tag": "Oliy Ta'lim", "title": "Oliy Ta'lim Qamrovi", "accent_color": ACCENT, "items": ["Oliy ta'lim bilan qamrovni 9% dan 38% ga oshirish.", "Nufuzli xorijiy universitetlar filiallarini ochish.", "OTMlarga akademik va moliyaviy mustaqillik berilishi."]}
    ]
    slide10 = prs.slides.add_slide(blank_layout)
    create_cards_layout(slide10, "4. Ta'lim Sohasidagi Tub O'zgarishlar va Inson Kapitali", "REJA 4: TA'LIM ISLOHOTLARI", cards10)
    add_footer(slide10, 10)

    # SLIDE 11
    cards11 = [
        {"title": "Tibbiy Xizmat Sifati", "accent_color": PRIMARY, "items": ["Birinchi bo'g'in (oilaviy shifokorlik va poliklinikalar)ni rivojlantirish.", "Qishloq joylarda tibbiy maskanlar va tez yordam sifatini oshirish.", "Tibbiyot xodimlarining moddiy va ma'naviy rag'batlantirilishi."]},
        {"title": "Sug'urta va Shoshilinch Yordam", "accent_color": SECONDARY, "items": ["Davlat tibbiy sug'urtasi tizimini bosqichma-bosqich joriy etish.", "Yuqori texnologik murakkab operatsiyalarni mahalliy darajada o'tkazish.", "Dori-darmon ta'minotida shaffoflik va nazoratni kuchaytirish."]},
        {"title": "Aholi Salomatligi", "accent_color": ACCENT, "items": ["Ommaviy sport va sog'lom turmush tarzini targ'ib qilish.", "Ona va bola salomatligini muhofaza qilish bo'yicha milliy dasturlar.", "Aholining o'rtacha umr ko'rish davomiyligining uzayishi."]}
    ]
    slide11 = prs.slides.add_slide(blank_layout)
    create_cards_layout(slide11, "Sog'liqni Saqlash Tizimi va Aholi Salomatligi Muhofazasi", "REJA 4: SOG'LIQNI SAQLASH", cards11)
    add_footer(slide11, 11)

    # SLIDE 12
    cards12 = [
        {"tag": "Prioritet", "title": "Markaziy Osiyo - Asosiy Ustuvorlik", "accent_color": PRIMARY, "items": ["Qo'shni davlatlar bilan do'stona va yaxshi qo'shnichilik aloqalarini tiklash.", "Chegara va suv-energetika muammolarini tinch muloqot orqali hal etish.", "Mintaqaviy savdo-sotiq va transport koridorlarini rivojlantirish."]},
        {"tag": "Mintaqaviy Muloqot", "title": "Mintaqaviy Integratsiya", "accent_color": SECONDARY, "items": ["Markaziy Osiyo davlatlari rahbarlarining Maslahat uchrashuvlari tashabbusi.", "Mintaqada o'zaro ishonch va barqarorlik muhitining qaror topishi.", "Mintaqaviy xavfsizlik va birdamlikni ta'minlash."]},
        {"tag": "Ko'p Tomonlama", "title": "Xalqaro Tashkilotlar", "accent_color": ACCENT, "items": ["BMT, ShHT, TDT va MDB doirasida faol va tashabbuskor diplomatiya.", "BMT Bosh Assambleyasida global va mintaqaviy tashabbuslarning ilgari surilishi.", "O'zbekistonning xalqaro sammitlar markaziga aylanishi."]}
    ]
    slide12 = prs.slides.add_slide(blank_layout)
    create_cards_layout(slide12, "5. Tashqi Siyosat: Markaziy Osiyodagi Yangi Muhit", "REJA 5: TASHQI SIYOSAT", cards12)
    add_footer(slide12, 12)

    # SLIDE 13
    cards13 = [
        {"title": "Transport va Logistika", "accent_color": PRIMARY, "items": ["Xitoy - Qirg'iziston - O'zbekiston temir yo'li loyihasi.", "Trans-Afg'on transport koridori va Janubiy Osiyoga chiqish imkoniyati.", "Yangi strategik avto va temir yo'l yo'nalishlarini barpo etish."]},
        {"title": "Ekologik Diplomatiya", "accent_color": SECONDARY, "items": ["Orolbo'yi mintaqasi bo'yicha global tashabbuslar va BMT rezolutsiyalari.", "'Yashil makon' umummilliy loyihasi va ekologik barqarorlik.", "Muqobil va yashil energetikaga o'tish bo'yicha xalqaro hamkorlik."]},
        {"title": "Ko'p Vektorli Diplomatiya", "accent_color": ACCENT, "items": ["AQSH, Yevropa Ittifoqi, Xitoy, Rossiya va Osiyo mamlakatlari bilan muvozanatli aloqalar.", "Madaniy-gumanitar va turizm sohasidagi hamkorlikni kengaytirish.", "O'zbekistonning iqtisodiy diplomatiyasi va eksport geografiyasini oshirish."]}
    ]
    slide13 = prs.slides.add_slide(blank_layout)
    create_cards_layout(slide13, "Global Tashabbuslar, Logistika va Ekologik Diplomatiya", "REJA 5: GLOBAL TASHABBUSLAR", cards13)
    add_footer(slide13, 13)

    # SLIDE 14
    cards14 = [
        {"tag": "E'tirof 1", "title": "BMT va Xalqaro Tashkilotlar", "accent_color": PRIMARY, "items": ["O'zbekiston ilgari surgan bir necha maxsus rezolutsiyalarning BMT tomonidan qabul qilinishi.", "Inson huquqlari bo'yicha Kengashga a'zolik va faoliyat.", "YuNESKO va boshqa xalqaro tuzilmalar bilan samarali sheriklik."]},
        {"tag": "E'tirof 2", "title": "Nufuzli Nashrlar va Reytinglar", "accent_color": SECONDARY, "items": ["'The Economist' tomonidan O'zbekistonning 'Yil mamlakati' deb e'tirof etilishi.", "Jahon bankining 'Doing Business' va BMTning inson taraqqiyoti indeksidagi ko'tarilish.", "Sayyohlik va turizm jozibadorligi bo'yicha yuqori o'rinlar."]},
        {"tag": "E'tirof 3", "title": "Mintaqaviy Liderlik", "accent_color": ACCENT, "items": ["Markaziy Osiyoda tinchlik va hamkorlik kafolatchisi sifatida e'tirof etilishi.", "Afg'onistonda tinchlik o'rnatish va gumanitar yordam ko'rsatishdagi yetakchi o'rni.", "Mintaqada barqaror rivojlanish drayveriga aylanishi."]}
    ]
    slide14 = prs.slides.add_slide(blank_layout)
    create_cards_layout(slide14, "Yangi O'zbekiston Islohotlarining Xalqaro E'tirofi", "XALQARO E'TIROF", cards14)
    add_footer(slide14, 14)

    # SLIDE 15
    cards15 = [
        {"tag": "Xulosa 1", "title": "Taraqqiyotning Ortga Qaytmasligi", "accent_color": PRIMARY, "items": ["Yangi O'zbekiston strategiyasi - mamlakatimiz taraqqiyotining yangi poydevorini yaratdi.", "Islohotlar ortga qaytmas tus oldi va xalq tomonidan to'liq qo'llab-quvvatlanmoqda.", "Davlat va jamiyat o'rtasidagi ishonch ko'prigi mustahkamlandi."]},
        {"tag": "Xulosa 2", "title": "Inson Qadri - Doimiy Maydon", "accent_color": SECONDARY, "items": ["Inson qadrini ulug'lash - faqat shior emas, amaliy harakatlar mezoniga aylandi.", "Har bir fuqaro uchun teng imkoniyatlar va adolatli jamiyat barpo etilmoqda.", "Kelajak avlod uchun kuchli, mustaqil va farovon davlat meros qoladi."]},
        {"tag": "Xulosa 3", "title": "Kelajakka Nigoh", "accent_color": ACCENT, "items": ["O'zbekiston-2030 strategiyasi doirasida yanada yuksak marralar belgilandi.", "Raqamli iqtisodiyot, yashil energetika va innovatsion ta'lim ustuvor bo'lib qoladi.", "Milliy yuksalish va farovonlik sari sabit qadam tashlanmoqda."]}
    ]
    slide15 = prs.slides.add_slide(blank_layout)
    create_cards_layout(slide15, "Xulosa: Yangi O'zbekistonning Tarixiy Roli va Kelajagi", "XULOSA VA ISTIQBOL", cards15)
    add_footer(slide15, 15)

    # SLIDE 16
    slide16 = prs.slides.add_slide(blank_layout)
    bg16 = slide16.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg16.fill.solid()
    bg16.fill.fore_color.rgb = PRIMARY
    bg16.line.fill.background()

    tb16 = slide16.shapes.add_textbox(Inches(1.5), Inches(2.2), Inches(10.333), Inches(3.5))
    tf16 = tb16.text_frame
    tf16.word_wrap = True

    p16_1 = tf16.paragraphs[0]
    p16_1.text = "E'TIBORINGIZ UCHUN RAHMAT!"
    p16_1.font.size = Pt(38)
    p16_1.font.bold = True
    p16_1.font.color.rgb = ACCENT
    p16_1.alignment = PP_ALIGN.CENTER

    p16_2 = tf16.add_paragraph()
    p16_2.text = "Yangi O'zbekiston Taraqqiyot Strategiyasi va Inson Qadri Tamoyili"
    p16_2.font.size = Pt(18)
    p16_2.font.color.rgb = WHITE
    p16_2.alignment = PP_ALIGN.CENTER
    p16_2.space_before = Pt(15)

    p16_3 = tf16.add_paragraph()
    p16_3.text = "Savollar va muhokama uchun tashakkur!"
    p16_3.font.size = Pt(14)
    p16_3.font.color.rgb = RGBColor(200, 220, 245)
    p16_3.alignment = PP_ALIGN.CENTER
    p16_3.space_before = Pt(30)

    out_dir = r"C:\Users\user\Desktop\SlideTranslate_AI\generated_presentations"
    os.makedirs(out_dir, exist_ok=True)
    out_path = os.path.join(out_dir, "Yangi_Ozbekiston_Taraqqiyot_Strategiyasi.pptx")
    prs.save(out_path)
    print("SLIDE_BUILD_SUCCESS: " + out_path)

if __name__ == '__main__':
    create_presentation()
