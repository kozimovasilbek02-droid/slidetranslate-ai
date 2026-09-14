# -*- coding: utf-8 -*-
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def build_deck():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    NAVY = RGBColor(11, 37, 89)
    BLUE = RGBColor(0, 102, 204)
    EMERALD = RGBColor(16, 185, 129)
    GOLD = RGBColor(245, 158, 11)
    BG_LIGHT = RGBColor(248, 250, 252)
    WHITE = RGBColor(255, 255, 255)
    TEXT_DARK = RGBColor(15, 23, 42)
    TEXT_MUTED = RGBColor(100, 116, 139)
    BORDER_LIGHT = RGBColor(226, 232, 240)
    TOTAL_SLIDES = 16

    def apply_bg(slide, color=BG_LIGHT):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = color
        bg.line.fill.background()
        slide.shapes._spTree.remove(bg._element)
        slide.shapes._spTree.insert(2, bg._element)

    def add_header(slide, title, category="O'ZBEKISTONNING ENG YANGI TARIXI"):
        badge = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.4), Inches(3.4), Inches(0.35))
        badge.fill.solid()
        badge.fill.fore_color.rgb = RGBColor(238, 242, 255)
        badge.line.color.rgb = BLUE
        badge.line.width = Pt(1)
        tf_b = badge.text_frame
        p_b = tf_b.paragraphs[0]
        p_b.text = f"✦  {category.upper()}"
        p_b.font.size = Pt(9.5)
        p_b.font.bold = True
        p_b.font.color.rgb = BLUE
        p_b.alignment = PP_ALIGN.CENTER

        tb = slide.shapes.add_textbox(Inches(0.8), Inches(0.78), Inches(11.733), Inches(0.75))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(24)
        p.font.bold = True
        p.font.color.rgb = NAVY

    def add_footer(slide, current):
        line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(6.85), Inches(11.733), Inches(0.015))
        line.fill.solid()
        line.fill.fore_color.rgb = BORDER_LIGHT
        line.line.fill.background()

        tb = slide.shapes.add_textbox(Inches(0.8), Inches(6.92), Inches(11.733), Inches(0.35))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.text = f"Yangi O'zbekiston Taraqqiyot Strategiyasi va Inson Qadri Tamoyili  |  Slayd {current} / {TOTAL_SLIDES}"
        p.font.size = Pt(9.5)
        p.font.color.rgb = TEXT_MUTED

    # SLIDE 1: COVER
    slide1 = prs.slides.add_slide(blank_layout)
    apply_bg(slide1, NAVY)

    accent_circle = slide1.shapes.add_shape(MSO_SHAPE.OVAL, Inches(9.5), Inches(-1.5), Inches(6.0), Inches(6.0))
    accent_circle.fill.solid()
    accent_circle.fill.fore_color.rgb = RGBColor(16, 50, 115)
    accent_circle.line.fill.background()

    tag_box = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(1.2), Inches(4.0), Inches(0.42))
    tag_box.fill.solid()
    tag_box.fill.fore_color.rgb = RGBColor(254, 243, 199)
    tag_box.line.color.rgb = GOLD
    tf_tag = tag_box.text_frame
    p_tag = tf_tag.paragraphs[0]
    p_tag.text = "✦  O'ZBEKISTONNING ENG YANGI TARIXI"
    p_tag.font.size = Pt(10.5)
    p_tag.font.bold = True
    p_tag.font.color.rgb = RGBColor(180, 83, 9)
    p_tag.alignment = PP_ALIGN.CENTER

    tb_h = slide1.shapes.add_textbox(Inches(1.0), Inches(1.9), Inches(11.3), Inches(4.5))
    tf_h = tb_h.text_frame
    tf_h.word_wrap = True
    
    p1 = tf_h.paragraphs[0]
    p1.text = "Yangi O'zbekiston Taraqqiyot Strategiyasi va Inson Qadri Tamoyili"
    p1.font.size = Pt(36)
    p1.font.bold = True
    p1.font.color.rgb = WHITE
    
    p2 = tf_h.add_paragraph()
    p2.text = "Inson qadrini ulug'lash, erkin fuqarolik jamiyatini barpo etish va milliy iqtisodiyotni yuksaltirishga qaratilgan strategik islohotlarning atroflicha ilmiy tahlili"
    p2.font.size = Pt(16)
    p2.font.color.rgb = RGBColor(203, 213, 225)
    p2.space_before = Pt(20)

    stats = [
        ("16 Slayd", "Kengaytirilgan visual format"),
        ("Tarix & Siyosat", "Fanlararo ilmiy yondashuv"),
        ("2022 - 2026", "Taraqqiyot strategiyasi")
    ]
    for idx, (st_t, st_d) in enumerate(stats):
        x = Inches(1.0) + idx * Inches(3.8)
        sb = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(5.6), Inches(3.5), Inches(1.0))
        sb.fill.solid()
        sb.fill.fore_color.rgb = RGBColor(20, 50, 110)
        sb.line.color.rgb = RGBColor(40, 85, 170)
        tf_s = sb.text_frame
        tf_s.word_wrap = True
        p_st = tf_s.paragraphs[0]
        p_st.text = st_t
        p_st.font.size = Pt(15)
        p_st.font.bold = True
        p_st.font.color.rgb = GOLD
        p_sd = tf_s.add_paragraph()
        p_sd.text = st_d
        p_sd.font.size = Pt(10)
        p_sd.font.color.rgb = RGBColor(203, 213, 225)

    # SLIDE 2: MUNDARIJA
    slide2 = prs.slides.add_slide(blank_layout)
    apply_bg(slide2)
    add_header(slide2, "Taqdimot Rejasi va Tuzilishi (Mundarija)", "MUNDARIJA")
    add_footer(slide2, 2)

    agenda = [
        ("01", "Yangi O'zbekiston Tushunchasi va Strategik Maqsadlar", "Milliy tiklanishdan milliy yuksalish sari o'tish bosqichlari"),
        ("02", "Inson Qadri va Uning Davlat Siyosatidagi Ustuvor O'rni", "Inson - jamiyat - davlat tamoyili va ijtimoiy kafolatlar"),
        ("03", "Iqtisodiy Islohotlar: Erkin Bozor va Investitsion Muhit", "Tadbirkorlikni qo'llab-quvvatlash va raqamli iqtisodiyot"),
        ("04", "Ta'lim va Sog'liqni Saqlash Sohasidagi Tub O'zgarishlar", "Inson kapitaliga investitsiya va tibbiy xizmatlar sifati"),
        ("05", "Tashqi Siyosat va Markaziy Osiyodagi Yangi Muhit", "Ochoq, pragmatik va yaxshi qo'shnichilik diplomatiyasi")
    ]

    for idx, (num, title, desc) in enumerate(agenda):
        y = Inches(1.65) + idx * Inches(0.98)
        card = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), y, Inches(11.733), Inches(0.88))
        card.fill.solid()
        card.fill.fore_color.rgb = WHITE
        card.line.color.rgb = BORDER_LIGHT
        card.line.width = Pt(1)

        num_badge = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), y + Inches(0.14), Inches(0.6), Inches(0.6))
        num_badge.fill.solid()
        num_badge.fill.fore_color.rgb = BLUE
        num_badge.line.fill.background()
        tf_n = num_badge.text_frame
        p_n = tf_n.paragraphs[0]
        p_n.text = num
        p_n.font.size = Pt(14)
        p_n.font.bold = True
        p_n.font.color.rgb = WHITE
        p_n.alignment = PP_ALIGN.CENTER

        tb = slide2.shapes.add_textbox(Inches(1.8), y + Inches(0.12), Inches(10.5), Inches(0.65))
        tf = tb.text_frame
        tf.word_wrap = True
        p_t = tf.paragraphs[0]
        p_t.text = title
        p_t.font.size = Pt(14)
        p_t.font.bold = True
        p_t.font.color.rgb = NAVY
        
        p_d = tf.add_paragraph()
        p_d.text = desc
        p_d.font.size = Pt(10.5)
        p_d.font.color.rgb = TEXT_MUTED

    def build_3card_slide(prs, title, category, slide_num, cards_data):
        slide = prs.slides.add_slide(blank_layout)
        apply_bg(slide)
        add_header(slide, title, category)
        add_footer(slide, slide_num)

        for idx, cd in enumerate(cards_data):
            x = Inches(0.8) + idx * Inches(3.98)
            y = Inches(1.65)
            w = Inches(3.77)
            h = Inches(4.9)

            card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
            card.fill.solid()
            card.fill.fore_color.rgb = WHITE
            card.line.color.rgb = BORDER_LIGHT
            card.line.width = Pt(1)

            hbar = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, Inches(0.85))
            hbar.fill.solid()
            hbar.fill.fore_color.rgb = cd.get("bg_color", RGBColor(241, 245, 249))
            hbar.line.fill.background()

            tb_h = slide.shapes.add_textbox(x + Inches(0.2), y + Inches(0.12), w - Inches(0.4), Inches(0.65))
            tf_h = tb_h.text_frame
            tf_h.word_wrap = True
            p_tag = tf_h.paragraphs[0]
            p_tag.text = cd.get("tag", "BO'LIM").upper()
            p_tag.font.size = Pt(9)
            p_tag.font.bold = True
            p_tag.font.color.rgb = cd.get("accent_color", BLUE)
            
            p_ct = tf_h.add_paragraph()
            p_ct.text = cd["title"]
            p_ct.font.size = Pt(13.5)
            p_ct.font.bold = True
            p_ct.font.color.rgb = NAVY

            tb_b = slide.shapes.add_textbox(x + Inches(0.2), y + Inches(1.0), w - Inches(0.4), h - Inches(1.1))
            tf_b = tb_b.text_frame
            tf_b.word_wrap = True

            for p_idx, pt_text in enumerate(cd["items"]):
                p_item = tf_b.paragraphs[0] if p_idx == 0 else tf_b.add_paragraph()
                p_item.text = f"✔  {pt_text}"
                p_item.font.size = Pt(11)
                p_item.font.color.rgb = TEXT_DARK
                p_item.space_before = Pt(8) if p_idx > 0 else Pt(0)

    slides_content = [
        (3, "Yangi Davrning Tarixiy Zaruriyati va Bosh G'oyasi", "KIRISH VA KONSEPSIYA", [
            {"tag": "Tarixiy Kontekst", "title": "Tarixiy Burilish Bosqichi", "accent_color": BLUE, "bg_color": RGBColor(238, 242, 255), "items": ["2016-yildan e'tiboran O'zbekiston rivojlanishining mutlaqo yangi davri boshlandi.", "Tarixiy taraqqiyotda davlat va jamiyat o'rtasidagi munosabatlar tubdan o'zgardi.", "Eski byurokratik g'ovlar olib tashlanib, islohotlarning ochiqligi ta'minlandi."]},
            {"tag": "Konsepsiya", "title": "Tarixiy Konsepsiya", "accent_color": EMERALD, "bg_color": RGBColor(236, 253, 245), "items": ["'Milliy tiklanishdan - Milliy yuksalish sari' g'oyasi bosh tamoyilga aylandi.", "Islohotlar faqat raqamlar uchun emas, balki har bir fuqaro hayotida sezilishi shart qilindi.", "Tarixiy tajriba va zamonaviy xalqaro standartlar uyg'unlashtirildi."]},
            {"tag": "Maqsad", "title": "Bosh Maqsad va Vazifa", "accent_color": GOLD, "bg_color": RGBColor(254, 243, 199), "items": ["Erkin, obod va farovon Yangi O'zbekistonni barpo etish.", "Dunyo hamjamiyatida mamlakatning munosib o'rni va nufuzini ta'minlash.", "Kelajak avlod uchun barqaror rivojlanish poydevorini yaratish."]}
        ]),
        (4, "1. Yangi O'zbekiston Taraqqiyot Strategiyasi (2022-2026)", "REJA 1: STRATEGIK MAQSADLAR", [
            {"tag": "Yo'nalishlar", "title": "7 Ta Ustuvor Yo'nalish", "accent_color": BLUE, "bg_color": RGBColor(238, 242, 255), "items": ["Erkin fuqarolik jamiyatini rivojlantirish.", "Adolat va qonun ustuvorligini ta'minlash.", "Milliy iqtisodiyotni jadal rivojlantirish.", "Adolatli ijtimoiy siyosat yuritish."]},
            {"tag": "Maqsadlar", "title": "Ma'naviyat va Xavfsizlik", "accent_color": EMERALD, "bg_color": RGBColor(236, 253, 245), "items": ["Ma'naviy taraqqiyotni ta'minlash va sohani yangi bosqichga olib chiqish.", "Milliy manfaatlardan kelib chiqqan holda tashqi siyosat yuritish.", "Mamlakat xavfsizligi va mudofaa salohiyatini kuchaytirish."]},
            {"tag": "Indikatorlar", "title": "Kutilayotgan Natijalar", "accent_color": GOLD, "bg_color": RGBColor(254, 243, 199), "items": ["Aholi jon boshiga yalpi ichki mahsulot hajmini oshirish.", "O'rta daromadli mamlakatlar qatoridan munosib o'rin egallash.", "Kambag'allik darajasini kamida 2 baravarga qisqartirish."]}
        ]),
        (5, "Davlat Boshqaruvi va Mahalla Tizimida Transformatsiya", "REJA 1: DAVLAT BOSHQARUVI", [
            {"tag": "Boshqaruv", "title": "Tizimli Transformatsiya", "accent_color": BLUE, "bg_color": RGBColor(238, 242, 255), "items": ["Ixcham va samarali davlat boshqaruv apparatini shakllantirish.", "Vazirlik va idoralar mas'uliyatini oshirish va maqbullashtirish.", "Davlat xizmatlarini to'liq raqamlashtirish va shaffoflikni ta'minlash."]},
            {"tag": "Mahalla", "title": "Mahallabay Tizimi", "accent_color": EMERALD, "bg_color": RGBColor(236, 253, 245), "items": ["Mahalla - jamiyatning tayanch bo'g'iniga aylantirildi.", "'Hokim yordamchisi', 'Yoshlar yetakchisi' va 'Xotin-qizlar faoli' institutlari joriy etildi.", "Muammolarni bevosita joyida hal etish tizimi yo'lga qo'yildi."]},
            {"tag": "Muloqot", "title": "Xalq Bilan Muloqot", "accent_color": GOLD, "bg_color": RGBColor(254, 243, 199), "items": ["Prezident Xalq qabulxonalari faoliyatining samaradorligi.", "Elektron hukumat (my.gov.uz) orqali yuzlab xizmatlarning yo'lga qo'yilishi.", "Jamoatchilik nazorati va fuqarolar tashabbusining oshishi."]}
        ]),
        (6, "2. 'Inson Qadri Uchun' Tamoyilining Mazmun-Mohiyati", "REJA 2: INSON QADRI", [
            {"tag": "Tamoyil 1", "title": "Inson Qadri - Bosh Mezon", "accent_color": BLUE, "bg_color": RGBColor(238, 242, 255), "items": ["'Inson - jamiyat - davlat' tamoyili ustuvor etib belgilandi.", "Davlat idoralari xalqqa xizmat qilishi shartligi konstitutsiyaviy darajada mustahkamlandi.", "Har bir fuqaro huquqi va erkinligi oliy qadriyat sanaladi."]},
            {"tag": "Tamoyil 2", "title": "Ijtimoiy Adolat va Himoya", "accent_color": EMERALD, "bg_color": RGBColor(236, 253, 245), "items": ["Kam ta'minlangan va ehtiyojmand oilalarni manzilli qo'llab-quvvatlash.", "'Ayollar daftari', 'Yoshlar daftari' va 'Saxovat daftari' tizimlari samaradorligi.", "Nogironligi bo'lgan shaxslarga mehir va e'tibor qaratish."]},
            {"tag": "Tamoyil 3", "title": "Konstitutsiyaviy Islohotlar", "accent_color": GOLD, "bg_color": RGBColor(254, 243, 199), "items": ["Yangi tahrirdagi Konstitutsiyada 'Ijtimoiy Davlat' maqomining belgilanishi.", "Insonning mehnati, yashashi va ta'lim olishiga bo'lgan huquqlarining kengaytirilishi.", "Shaxsiy daxlsizlik va xususiy mulk kafolatlarining kuchaytirilishi."]}
        ]),
        (7, "Inson Huquqlari, Sud Mustaqilligi va So'z Erkinligi", "REJA 2: HUQUQIY KAFOLATLAR", [
            {"tag": "Sud-Huquq", "title": "Sud Mustaqilligi", "accent_color": BLUE, "bg_color": RGBColor(238, 242, 255), "items": ["Sudlarning haqiqiy mustaqilligini ta'minlash.", "Inson huquqlari poymol etuvchi har qanday harakatga barham berish.", "Advokatura institutini va himoya huquqini kuchaytirish."]},
            {"tag": "Mehnat", "title": "Majburiy Mehnatga Barham", "accent_color": EMERALD, "bg_color": RGBColor(236, 253, 245), "items": ["Bolalar mehnati va majburiy mehnat to'liq tugatildi.", "Xalqaro Mehnat Tashkiloti (XMT) tomonidan e'tirof etildi.", "Paxta boykoti (Cotton Campaign) bekor qilinishiga erishildi."]},
            {"tag": "Erkinlik", "title": "So'z Erkinligi va OAV", "accent_color": GOLD, "bg_color": RGBColor(254, 243, 199), "items": ["Ommaviy axborot vositalari (OAV) va blogerlar faoliyatiga keng imkoniyatlar.", "Davlat organlari faoliyatining ochiqligi va shaffofligi.", "Jamoatchilik fikrining davlat qarorlariga ta'sir o'tkazishi."]}
        ]),
        (8, "3. Iqtisodiy Islohotlar: Valyuta, Soliq va Erkin Bozor", "REJA 3: IQTISODIYOT", [
            {"tag": "Valyuta", "title": "Valyuta Bozorini Erkinlashtirish", "accent_color": BLUE, "bg_color": RGBColor(238, 242, 255), "items": ["2017-yildagi erkin valyuta konvertatsiyasining joriy etilishi.", "Xorijiy investorlar va mahalliy tadbirkorlar uchun teng sharoitlar yaratish.", "Valyuta cheklovlarining bekor qilinishi."]},
            {"tag": "Soliqlar", "title": "Soliq Yukini Kamaytirish", "accent_color": EMERALD, "bg_color": RGBColor(236, 253, 245), "items": ["Soliq stavkalarining maqbullashtirilishi va soddalashtirilishi.", "QQS (Qo'shilgan qiymat solig'i) stavkasining 15% dan 12% ga tushirilishi.", "Halol tadbirkorlik uchun imtiyozli sharoitlar."]},
            {"tag": "Bozor", "title": "Mulkchilik va Xususiylashtirish", "accent_color": GOLD, "bg_color": RGBColor(254, 243, 199), "items": ["Xususiy mulk daxlsizligining qonuniy kafolatlanishi.", "Davlat aktivlarini shaffof tenderlar orqali xususiylashtirish.", "Monopoliyani kamaytirish va raqobat muhitini rivojlantirish."]}
        ]),
        (9, "Investitsiyalar Jalb Etish va Tadbirkorlikni Qo'llab-quvvatlash", "REJA 3: INVESTITSIYA VA BIZNES", [
            {"tag": "Investitsiya", "title": "Investitsiya Muhiti", "accent_color": BLUE, "bg_color": RGBColor(238, 242, 255), "items": ["Xorijiy sarmoyadorlar uchun huquqiy va iqtisodiy kafolatlar.", "Maxsus iqtisodiy zonalar (MIZ) va texnoparklar tarmog'ini kengaytirish.", "Xalqaro kredit reytinqlarida O'zbekiston pozitsiyasining yaxshilanishi."]},
            {"tag": "Biznes", "title": "Kichik va O'rta Biznes", "accent_color": EMERALD, "bg_color": RGBColor(236, 253, 245), "items": ["Tadbirkorlik faoliyatini tekshirishlarga moratoriy va yengilliklar.", "Imtiyozli kreditlar va subsidiyalar ajratish hajmining oshishi.", "Yangi ish o'rinlarini yaratishda biznesning asosiy drayverga aylanishi."]},
            {"tag": "Sanoat", "title": "Sanoat va Eksport", "accent_color": GOLD, "bg_color": RGBColor(254, 243, 199), "items": ["Xomashyo eksportidan tayyor mahsulot eksportiga o'tish strategiyasi.", "Avtomobilsozlik, to'qimachilik va kimyo sanoatida modernizatsiya.", "Jahon Savdo Tashkilotiga (JST) a'zo bo'lish bo'yicha jadal muzokaralar."]}
        ]),
        (10, "4. Ta'lim Sohasidagi Tub O'zgarishlar va Inson Kapitali", "REJA 4: TA'LIM ISLOHOTLARI", [
            {"tag": "Bog'cha", "title": "Maktabgacha Ta'lim", "accent_color": BLUE, "bg_color": RGBColor(238, 242, 255), "items": ["Maktabgacha ta'lim vazirligining tashkil etilishi.", "Qamrov darajasini 27% dan 72% dan ortiq ko'rsatkichga etkazish.", "Xususiy va davlat-xususiy sheriklik bog'chalarini ko'paytirish."]},
            {"tag": "Maktab", "title": "Maktablar va O'qituvchi Qadri", "accent_color": EMERALD, "bg_color": RGBColor(236, 253, 245), "items": ["O'qituvchilar ish haqini oshirish va ularni majburiy mehnatsiz qilish.", "Prezident maktablari va Ijod maktablari tarmog'ini yaratish.", "Zamonaviy darsliklar va milliy o'quv dasturlarini joriy etish."]},
            {"tag": "OTM", "title": "Oliy Ta'lim Qamrovi", "accent_color": GOLD, "bg_color": RGBColor(254, 243, 199), "items": ["Oliy ta'lim bilan qamrovni 9% dan 38% ga oshirish.", "Nufuzli xorijiy universitetlar filiallarini ochish.", "OTMlarga akademik va moliyaviy mustaqillik berilishi."]}
        ]),
        (11, "Sog'liqni Saqlash Tizimi va Aholi Salomatligi Muhofazasi", "REJA 4: SOG'LIQNI SAQLASH", [
            {"tag": "Tibbiyot", "title": "Tibbiy Xizmat Sifati", "accent_color": BLUE, "bg_color": RGBColor(238, 242, 255), "items": ["Birinchi bo'g'in (oilaviy shifokorlik va poliklinikalar)ni rivojlantirish.", "Qishloq joylarda tibbiy maskanlar va tez yordam sifatini oshirish.", "Tibbiyot xodimlarining moddiy va ma'naviy rag'batlantirilishi."]},
            {"tag": "Sug'urta", "title": "Sug'urta va Texnologiya", "accent_color": EMERALD, "bg_color": RGBColor(236, 253, 245), "items": ["Davlat tibbiy sug'urtasi tizimini bosqichma-bosqich joriy etish.", "Yuqori texnologik murakkab operatsiyalarni mahalliy darajada o'tkazish.", "Dori-darmon ta'minotida shaffoflik va nazoratni kuchaytirish."]},
            {"tag": "Salomatlik", "title": "Aholi Salomatligi Muhofazasi", "accent_color": GOLD, "bg_color": RGBColor(254, 243, 199), "items": ["Ommaviy sport va sog'lom turmush tarzini targ'ib qilish.", "Ona va bola salomatligini muhofaza qilish bo'yicha milliy dasturlar.", "Aholining o'rtacha umr ko'rish davomiyligining uzayishi."]}
        ]),
        (12, "5. Tashqi Siyosat: Markaziy Osiyodagi Yangi Muhit", "REJA 5: TASHQI SIYOSAT", [
            {"tag": "Mintaqa", "title": "Markaziy Osiyo Ustuvorligi", "accent_color": BLUE, "bg_color": RGBColor(238, 242, 255), "items": ["Qo'shni davlatlar bilan do'stona va yaxshi qo'shnichilik aloqalarini tiklash.", "Chegara va suv-energetika muammolarini tinch muloqot orqali hal etish.", "Mintaqaviy savdo-sotiq va transport koridorlarini rivojlantirish."]},
            {"tag": "Integratsiya", "title": "Mintaqaviy Integratsiya", "accent_color": EMERALD, "bg_color": RGBColor(236, 253, 245), "items": ["Markaziy Osiyo davlatlari rahbarlarining Maslahat uchrashuvlari tashabbusi.", "Mintaqada o'zaro ishonch va barqarorlik muhitining qaror topishi.", "Mintaqaviy xavfsizlik va birdamlikni ta'minlash."]},
            {"tag": "Diplomatiya", "title": "Xalqaro Tashkilotlar", "accent_color": GOLD, "bg_color": RGBColor(254, 243, 199), "items": ["BMT, ShHT, TDT va MDB doirasida faol va tashabbuskor diplomatiya.", "BMT Bosh Assambleyasida global va mintaqaviy tashabbuslarning ilgari surilishi.", "O'zbekistonning xalqaro sammitlar markaziga aylanishi."]}
        ]),
        (13, "Global Tashabbuslar, Logistika va Ekologik Diplomatiya", "REJA 5: GLOBAL TASHABBUSLAR", [
            {"tag": "Logistika", "title": "Transport Koridorlari", "accent_color": BLUE, "bg_color": RGBColor(238, 242, 255), "items": ["Xitoy - Qirg'iziston - O'zbekiston temir yo'li loyihasi.", "Trans-Afg'on transport koridori va Janubiy Osiyoga chiqish imkoniyati.", "Yangi strategik avto va temir yo'l yo'nalishlarini barpo etish."]},
            {"tag": "Ekologiya", "title": "Ekologik Diplomatiya", "accent_color": EMERALD, "bg_color": RGBColor(236, 253, 245), "items": ["Orolbo'yi mintaqasi bo'yicha global tashabbuslar va BMT rezolutsiyalari.", "'Yashil makon' umummilliy loyihasi va ekologik barqarorlik.", "Muqobil va yashil energetikaga o'tish bo'yicha xalqaro hamkorlik."]},
            {"tag": "Ko'p Vektorli", "title": "Muvozanatli Diplomatiya", "accent_color": GOLD, "bg_color": RGBColor(254, 243, 199), "items": ["AQSH, Yevropa Ittifoqi, Xitoy va Rossiya bilan muvozanatli aloqalar.", "Madaniy-gumanitar va turizm sohasidagi hamkorlikni kengaytirish.", "O'zbekistonning iqtisodiy diplomatiyasi va eksport geografiyasini oshirish."]}
        ]),
        (14, "Yangi O'zbekiston Islohotlarining Xalqaro E'tirofi", "XALQARO E'TIROF", [
            {"tag": "E'tirof 1", "title": "BMT va Tashkilotlar", "accent_color": BLUE, "bg_color": RGBColor(238, 242, 255), "items": ["O'zbekiston ilgari surgan bir necha maxsus rezolutsiyalarning BMT tomonidan qabul qilinishi.", "Inson huquqlari bo'yicha Kengashga a'zolik va faoliyat.", "YuNESKO va boshqa xalqaro tuzilmalar bilan samarali sheriklik."]},
            {"tag": "E'tirof 2", "title": "Nufuzli Reytinglar", "accent_color": EMERALD, "bg_color": RGBColor(236, 253, 245), "items": ["'The Economist' tomonidan O'zbekistonning 'Yil mamlakati' deb e'tirof etilishi.", "Jahon bankining 'Doing Business' va BMTning inson taraqqiyoti indeksidagi ko'tarilish.", "Sayyohlik va turizm jozibadorligi bo'yicha yuqori o'rinlar."]},
            {"tag": "E'tirof 3", "title": "Mintaqaviy Liderlik", "accent_color": GOLD, "bg_color": RGBColor(254, 243, 199), "items": ["Markaziy Osiyoda tinchlik va hamkorlik kafolatchisi sifatida e'tirof etilishi.", "Afg'onistonda tinchlik o'rnatish va gumanitar yordam ko'rsatishdagi yetakchi o'rni.", "Mintaqada barqaror rivojlanish drayveriga aylanishi."]}
        ]),
        (15, "Xulosa: Yangi O'zbekistonning Tarixiy Roli va Kelajagi", "XULOSA VA ISTIQBOL", [
            {"tag": "Xulosa 1", "title": "Taraqqiyotning Ortga Qaytmasligi", "accent_color": BLUE, "bg_color": RGBColor(238, 242, 255), "items": ["Yangi O'zbekiston strategiyasi - mamlakatimiz taraqqiyotining yangi poydevorini yaratdi.", "Islohotlar ortga qaytmas tus oldi va xalq tomonidan to'liq qo'llab-quvvatlanmoqda.", "Davlat va jamiyat o'rtasidagi ishonch ko'prigi mustahkamlandi."]},
            {"tag": "Xulosa 2", "title": "Inson Qadri - Doimiy Maydon", "accent_color": EMERALD, "bg_color": RGBColor(236, 253, 245), "items": ["Inson qadrini ulug'lash - faqat shior emas, amaliy harakatlar mezoniga aylandi.", "Har bir fuqaro uchun teng imkoniyatlar va adolatli jamiyat barpo etilmoqda.", "Kelajak avlod uchun kuchli, mustaqil va farovon davlat meros qoladi."]},
            {"tag": "Xulosa 3", "title": "Kelajakka Nigoh", "accent_color": GOLD, "bg_color": RGBColor(254, 243, 199), "items": ["O'zbekiston-2030 strategiyasi doirasida yanada yuksak marralar belgilandi.", "Raqamli iqtisodiyot, yashil energetika va innovatsion ta'lim ustuvor bo'lib qoladi.", "Milliy yuksalish va farovonlik sari sabit qadam tashlanmoqda."]}
        ])
    ]

    for num, t, c, cd in slides_content:
        build_3card_slide(prs, t, c, num, cd)

    # SLIDE 16: END
    slide16 = prs.slides.add_slide(blank_layout)
    apply_bg(slide16, NAVY)
    tb16 = slide16.shapes.add_textbox(Inches(1.5), Inches(2.2), Inches(10.333), Inches(3.5))
    tf16 = tb16.text_frame
    tf16.word_wrap = True
    p16_1 = tf16.paragraphs[0]
    p16_1.text = "E'TIBORINGIZ UCHUN RAHMAT!"
    p16_1.font.size = Pt(38)
    p16_1.font.bold = True
    p16_1.font.color.rgb = GOLD
    p16_1.alignment = PP_ALIGN.CENTER
    p16_2 = tf16.add_paragraph()
    p16_2.text = "Yangi O'zbekiston Taraqqiyot Strategiyasi va Inson Qadri Tamoyili"
    p16_2.font.size = Pt(18)
    p16_2.font.color.rgb = WHITE
    p16_2.alignment = PP_ALIGN.CENTER
    p16_2.space_before = Pt(15)
    p16_3 = tf16.add_paragraph()
    p16_3.text = "Savollar va muhokama uchun tashakkur!"
    p16_3.font.size = Pt(14)
    p16_3.font.color.rgb = RGBColor(203, 213, 225)
    p16_3.alignment = PP_ALIGN.CENTER
    p16_3.space_before = Pt(30)

    out_dir = r"C:\Users\user\Desktop\SlideTranslate_AI\generated_presentations"
    os.makedirs(out_dir, exist_ok=True)
    out_path = os.path.join(out_dir, "Yangi_Ozbekiston_Taraqqiyot_Strategiyasi_UltraModern.pptx")
    prs.save(out_path)
    print(f"SUCCESS_BUILD_DECK_EXPORTED: {out_path}")

if __name__ == '__main__':
    build_deck()
