# -*- coding: utf-8 -*-
import os
import sys

# Ensure UTF-8 stdout on Windows
if sys.stdout.encoding != 'utf-8':
    try:
        sys.stdout.reconfigure(encoding='utf-8')
    except Exception:
        pass

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

from backend.core.transliteration import latin_to_cyrillic, cyrillic_to_latin, ensure_script
from backend.core.pptx_processor import PPTXProcessor
from backend.core.gemini_translator import GeminiTranslator

def run_tests():
    print("=== 1. Testing Transliteration Engine ===")
    sample_lat = "O'zbekiston kelajagi buyuk davlat. Sun'iy intellekt va axborot texnologiyalari."
    sample_cyr = latin_to_cyrillic(sample_lat)
    print(f"Latin: {sample_lat}")
    print(f"Cyrillic: {sample_cyr}")
    back_lat = cyrillic_to_latin(sample_cyr)
    print(f"Back to Latin: {back_lat}")
    assert "Ўзбекистон" in sample_cyr, "Cyrillic conversion failed"
    assert "интеллект" in sample_cyr, "Cyrillic conversion failed"

    print("\n=== 2. Creating Sample Multi-slide PPTX for Testing ===")
    test_pptx_dir = "test_artifacts"
    os.makedirs(test_pptx_dir, exist_ok=True)
    sample_pptx_path = os.path.join(test_pptx_dir, "sample_presentation.pptx")

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5) # 16:9
    blank_layout = prs.slide_layouts[6]

    # Slide 1: Title & Subtitle & Bullets
    slide1 = prs.slides.add_slide(blank_layout)
    txBox1 = slide1.shapes.add_textbox(Inches(1.0), Inches(0.8), Inches(11.333), Inches(1.5))
    tf1 = txBox1.text_frame
    p1 = tf1.paragraphs[0]
    p1.text = "Strategic Roadmap 2030"
    p1.font.bold = True
    p1.font.size = Pt(36)
    p1.font.color.rgb = RGBColor(30, 58, 138)

    p2 = tf1.add_paragraph()
    p2.text = "Accelerating Digital Transformation & Global Innovation"
    p2.font.size = Pt(20)
    p2.font.color.rgb = RGBColor(75, 85, 99)

    # Bullet Box
    txBox2 = slide1.shapes.add_textbox(Inches(1.0), Inches(2.8), Inches(5.5), Inches(3.5))
    tf2 = txBox2.text_frame
    bp1 = tf2.paragraphs[0]
    bp1.text = "• AI-driven workflow optimization"
    bp1.font.size = Pt(16)
    bp2 = tf2.add_paragraph()
    bp2.text = "• High-precision localization and translation"
    bp2.font.size = Pt(16)
    bp3 = tf2.add_paragraph()
    bp3.text = "• Seamless enterprise integration"
    bp3.font.size = Pt(16)

    # Slide 2: Table
    slide2 = prs.slides.add_slide(blank_layout)
    t_box = slide2.shapes.add_textbox(Inches(1.0), Inches(0.8), Inches(11.333), Inches(1.0))
    t_box.text_frame.paragraphs[0].text = "Quarterly Milestones & Target Deliverables"
    t_box.text_frame.paragraphs[0].font.size = Pt(28)
    t_box.text_frame.paragraphs[0].font.bold = True

    table_shape = slide2.shapes.add_table(3, 3, Inches(1.0), Inches(2.2), Inches(11.0), Inches(3.0))
    tbl = table_shape.table
    tbl.cell(0, 0).text_frame.paragraphs[0].text = "Phase"
    tbl.cell(0, 1).text_frame.paragraphs[0].text = "Objective"
    tbl.cell(0, 2).text_frame.paragraphs[0].text = "Expected Timeline"

    tbl.cell(1, 0).text_frame.paragraphs[0].text = "Q1 2026"
    tbl.cell(1, 1).text_frame.paragraphs[0].text = "Core Engine Architecture"
    tbl.cell(1, 2).text_frame.paragraphs[0].text = "March 31"

    tbl.cell(2, 0).text_frame.paragraphs[0].text = "Q2 2026"
    tbl.cell(2, 1).text_frame.paragraphs[0].text = "Full Platform Rollout"
    tbl.cell(2, 2).text_frame.paragraphs[0].text = "June 30"

    prs.save(sample_pptx_path)
    print(f"Created sample test PPTX at: {sample_pptx_path}")

    print("\n=== 3. Testing PPTX Extractor ===")
    extracted = PPTXProcessor.extract_presentation_data(sample_pptx_path)
    print(f"Slides count: {extracted['slides_count']}, Total items: {extracted['total_items']}")
    assert extracted["slides_count"] == 2, f"Expected 2 slides, got {extracted['slides_count']}"
    assert extracted["total_items"] >= 9, f"Expected at least 9 items, got {extracted['total_items']}"
    for s in extracted["slides"]:
        print(f"  Slide {s['slide_index']}: '{s['title']}' ({s['items_count']} items)")

    print("\n=== 4. Testing Translations Mapping & Export (Lotin) ===")
    translations_lat = {
        "s1_sh0_p0": "2030-yilgi Strategik Yo'l Xaritasi",
        "s1_sh0_p1": "Raqamli Transformatsiya va Global Innovatsiyalarni Jadallashtirish",
        "s1_sh1_p0": "• Sun'iy intellekt asosida ish jarayonlarini optimallashtirish",
        "s1_sh1_p1": "• Yuqori aniqlikdagi mahalliylashtirish va tarjima",
        "s1_sh1_p2": "• Korxona tizimlariga uzluksiz integratsiya",
        "s2_sh0_p0": "Choraklik Muhim Bosqichlar va Maqsadli Natijalar",
        "s2_sh1_tbl_r0_c0_p0": "Bosqich",
        "s2_sh1_tbl_r0_c1_p0": "Maqsad",
        "s2_sh1_tbl_r0_c2_p0": "Kutilayotgan Muddat",
        "s2_sh1_tbl_r1_c0_p0": "2026-yil 1-chorak",
        "s2_sh1_tbl_r1_c1_p0": "Asosiy Dvigatel Arxitekturasi",
        "s2_sh1_tbl_r1_c2_p0": "31-mart",
        "s2_sh1_tbl_r2_c0_p0": "2026-yil 2-chorak",
        "s2_sh1_tbl_r2_c1_p0": "To'liq Platformani Ishga Tushirish",
        "s2_sh1_tbl_r2_c2_p0": "30-iyun",
    }

    out_lat_pptx = os.path.join(test_pptx_dir, "output_uzbek_latin.pptx")
    PPTXProcessor.apply_translations_and_export(
        original_pptx_path=sample_pptx_path,
        translations_map=translations_lat,
        output_pptx_path=out_lat_pptx,
        auto_fit=True,
        target_script="latin"
    )
    print(f"Generated Uzbek Latin PPTX: {out_lat_pptx}")
    assert os.path.exists(out_lat_pptx), "Latin PPTX was not created"

    # Verify extracted content of generated PPTX
    ver_lat = PPTXProcessor.extract_presentation_data(out_lat_pptx)
    assert "2030-yilgi Strategik" in ver_lat["slides"][0]["items"][0]["original_text"]
    print("Latin PPTX verification: PASSED!")

    print("\n=== 5. Testing Translations Mapping & Export (Kirill) ===")
    out_cyr_pptx = os.path.join(test_pptx_dir, "output_uzbek_cyrillic.pptx")
    PPTXProcessor.apply_translations_and_export(
        original_pptx_path=sample_pptx_path,
        translations_map=translations_lat,
        output_pptx_path=out_cyr_pptx,
        auto_fit=True,
        target_script="cyrillic"
    )
    print(f"Generated Uzbek Cyrillic PPTX: {out_cyr_pptx}")
    assert os.path.exists(out_cyr_pptx), "Cyrillic PPTX was not created"

    ver_cyr = PPTXProcessor.extract_presentation_data(out_cyr_pptx)
    assert "Стратегик" in ver_cyr["slides"][0]["items"][0]["original_text"]
    print("Cyrillic PPTX verification: PASSED!")

    print("\n=== 6. Testing Gemini Translator Module ===")
    translator = GeminiTranslator()
    mock_batch = [
        {"id": "test_1", "original_text": "Artificial Intelligence & Cloud Systems"},
        {"id": "test_2", "original_text": "Quarterly Financial Performance"}
    ]
    res_lat = translator.translate_items_batch(mock_batch, target_script="latin")
    res_cyr = translator.translate_items_batch(mock_batch, target_script="cyrillic")
    print("Translator result (Latin):", res_lat)
    print("Translator result (Cyrillic):", res_cyr)

    print("\n🎉 ALL BACKEND UNIT AND INTEGRATION TESTS PASSED SUCCESSFULLY!")

if __name__ == "__main__":
    run_tests()
