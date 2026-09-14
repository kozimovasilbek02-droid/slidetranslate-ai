import os
import asyncio
from playwright.async_api import async_playwright
from pptx import Presentation
from pptx.util import Inches

SLIDES_DATA = [
    {
        "bg": "linear-gradient(135deg, #0f172a 0%, #090d16 100%)",
        "badge": "O'ZBEKISTONNING ENG YANGI TARIXI",
        "title": "Yangi O'zbekiston Taraqqiyot Strategiyasi va Inson Qadri Tamoyili",
        "subtitle": "Inson qadrini ulug'lash, erkin fuqarolik jamiyatini barpo etish va milliy iqtisodiyotni yuksaltirishga qaratilgan strategik islohotlarning atroflicha ilmiy-vizual tahlili",
        "type": "hero"
    },
    {
        "bg": "linear-gradient(135deg, #0f172a 0%, #1e1b4b 100%)",
        "badge": "MUNDARIJA VA REJA",
        "title": "Taqdimot Rejasi va Asosiy Bo'limlar",
        "type": "agenda",
        "items": [
            ("01", "Yangi O'zbekiston Tushunchasi va Strategik Maqsadlar", "Milliy tiklanishdan milliy yuksalish sari o'tish bosqichlari"),
            ("02", "Inson Qadri va Uning Davlat Siyosatidagi Ustuvor O'rni", "Inson - jamiyat - davlat tamoyili va ijtimoiy kafolatlar"),
            ("03", "Iqtisodiy Islohotlar: Erkin Bozor va Investitsion Muhit", "Tadbirkorlikni qo'llab-quvvatlash va raqamli iqtisodiyot"),
            ("04", "Ta'lim va Sog'liqni Saqlash Sohasidagi Tub O'zgarishlar", "Inson kapitaliga investitsiya va tibbiy xizmatlar sifati"),
            ("05", "Tashqi Siyosat va Markaziy Osiyodagi Yangi Muhit", "Ochoq, pragmatik va yaxshi qo'shnichilik diplomatiyasi")
        ]
    },
    {
        "bg": "linear-gradient(135deg, #0f172a 0%, #172554 100%)",
        "badge": "KIRISH VA KONSEPSIYA",
        "title": "Yangi Davrning Tarixiy Zaruriyati va Bosh G'oyasi",
        "type": "cards",
        "cards": [
            ("Tarixiy Burilish Bosqichi", "2016-yildan e'tiboran O'zbekiston rivojlanishining mutlaqo yangi davri boshlandi.", "Davlat va jamiyat o'rtasidagi munosabatlar tubdan o'zgardi va ochiqlik ta'minlandi."),
            ("Tarixiy Konsepsiya", "'Milliy tiklanishdan - Milliy yuksalish sari' g'oyasi bosh tamoyilga aylandi.", "Islohotlar har bir fuqaro hayotida sezilishi shart qilib belgilandi."),
            ("Bosh Maqsad va Vazifa", "Erkin, obod va farovon Yangi O'zbekistonni barpo etish.", "Kelajak avlod uchun barqaror va qudratli rivojlanish poydevor yaratish.")
        ]
    },
    {
        "bg": "linear-gradient(135deg, #0f172a 0%, #064e3b 100%)",
        "badge": "REJA 1: STRATEGIK MAQSADLAR",
        "title": "1. Yangi O'zbekiston Taraqqiyot Strategiyasi (2022-2026)",
        "type": "cards",
        "cards": [
            ("7 Ta Ustuvor Yo'nalish", "Erkin fuqarolik jamiyati hamda adolat va qonun ustuvorligini ta'minlash.", "Milliy iqtisodiyotni jadal rivojlantirish va adolatli ijtimoiy siyosat."),
            ("Ma'naviyat va Xavfsizlik", "Ma'naviy taraqqiyotni ta'minlash va sohani yangi bosqichga olib chiqish.", "Mamlakat xavfsizligi va mudofaa salohiyatini keskin kuchaytirish."),
            ("Kutilayotgan Natijalar", "Aholi jon boshiga YIM hajmini oshirish va o'rta daromadli mamlakatlar safga kirish.", "Kambag'allik darajasini kamida 2 baravarga qisqartirish.")
        ]
    },
    {
        "bg": "linear-gradient(135deg, #0f172a 0%, #312e81 100%)",
        "badge": "REJA 1: DAVLAT BOSHQARUVI",
        "title": "Davlat Boshqaruvi va Mahalla Tizimida Transformatsiya",
        "type": "cards",
        "cards": [
            ("Tizimli Transformatsiya", "Ixcham va samarali davlat boshqaruv apparatini shakllantirish.", "Vazirliklar mas'uliyatini oshirish va xizmatlarni to'liq raqamlashtirish."),
            ("Mahallabay Tizimi", "Mahalla - jamiyatning tayanch bo'g'iniga aylantirildi.", "'Hokim yordamchisi', 'Yoshlar yetakchisi' va 'Xotin-qizlar faoli' institutlari."),
            ("Xalq Bilan Muloqot", "Prezident Xalq qabulxonalari va my.gov.uz elektron xizmatlari samaradorligi.", "Jamoatchilik nazorati va fuqarolar tashabbusining oshishi.")
        ]
    },
    {
        "bg": "linear-gradient(135deg, #0f172a 0%, #451a03 100%)",
        "badge": "REJA 2: INSON QADRI",
        "title": "2. 'Inson Qadri Uchun' Tamoyilining Mazmun-Mohiyati",
        "type": "cards",
        "cards": [
            ("Inson Qadri - Bosh Mezon", "'Inson - jamiyat - davlat' tamoyili ustuvor etib belgilandi.", "Davlat idoralari xalqqa xizmat qilishi konstitutsiyaviy majburiyat qilindi."),
            ("Ijtimoiy Adolat va Himoya", "Kam ta'minlangan oilalarni manzilli qo'llab-quvvatlash.", "'Ayollar daftari', 'Yoshlar daftari' va 'Saxovat daftari' tizimlari."),
            ("Konstitutsiyaviy Islohotlar", "Yangi tahrirdagi Konstitutsiyada 'Ijtimoiy Davlat' maqomining belgilanishi.", "Insonning mehnati va ta'lim olish huquqlarining kengaytirilishi.")
        ]
    },
    {
        "bg": "linear-gradient(135deg, #0f172a 0%, #1e293b 100%)",
        "badge": "REJA 2: HUQUQIY KAFOLATLAR",
        "title": "Inson Huquqlari, Sud Mustaqilligi va So'z Erkinligi",
        "type": "cards",
        "cards": [
            ("Sud Mustaqilligi", "Sudlarning haqiqiy mustaqilligini ta'minlash va inson huquqlarini himoya qilish.", "Advokatura institutini va himoya huquqini tubdan kuchaytirish."),
            ("Majburiy Mehnatga Barham", "Bolalar mehnati va majburiy mehnat to'liq tugatildi va XMT tomonidan e'tirof etildi.", "Paxta boykoti (Cotton Campaign) bekor qilinishiga erishildi."),
            ("So'z Erkinligi va OAV", "Ommaviy axborot vositalari va blogerlar faoliyatiga keng imkoniyatlar.", "Davlat organlari faoliyatining ochiqligi va shaffofligi.")
        ]
    },
    {
        "bg": "linear-gradient(135deg, #0f172a 0%, #065f46 100%)",
        "badge": "REJA 3: IQTISODIYOT",
        "title": "3. Iqtisodiy Islohotlar: Valyuta, Soliq va Erkin Bozor",
        "type": "cards",
        "cards": [
            ("Valyuta Bozorini Erkinlashtirish", "2017-yildagi erkin valyuta konvertatsiyasining joriy etilishi.", "Xorijiy investorlar va mahalliy biznes uchun teng sharoitlar yaratish."),
            ("Soliq Yukini Kamaytirish", "Soliq stavkalarining maqbullashtirilishi va soddalashtirilishi.", "QQS stavkasining 15% dan 12% ga tushirilishi hamda yengilliklar."),
            ("Mulkchilik va Xususiylashtirish", "Xususiy mulk daxlsizligining qonuniy kafolatlanishi.", "Davlat aktivlarini shaffof tenderlar orqali xususiylashtirish.")
        ]
    },
    {
        "bg": "linear-gradient(135deg, #0f172a 0%, #1e1b4b 100%)",
        "badge": "REJA 3: INVESTITSIYA VA BIZNES",
        "title": "Investitsiyalar Jalb Etish va Tadbirkorlikni Qo'llab-quvvatlash",
        "type": "cards",
        "cards": [
            ("Investitsiya Muhiti", "Xorijiy sarmoyadorlar uchun huquqiy va iqtisodiy kafolatlar.", "Maxsus iqtisodiy zonalar (MIZ) va texnoparklar tarmog'ini kengaytirish."),
            ("Kichik va O'rta Biznes", "Tadbirkorlik faoliyatini tekshirishlarga moratoriy va yengilliklar.", "Yangi ish o'rinlarini yaratishda biznesning drayverga aylanishi."),
            ("Sanoat va Eksport", "Xomashyo eksportidan tayyor mahsulot eksportiga o'tish strategiyasi.", "Jahon Savdo Tashkilotiga (JST) a'zo bo'lish bo'yicha jadal muzokaralar.")
        ]
    },
    {
        "bg": "linear-gradient(135deg, #0f172a 0%, #3730a3 100%)",
        "badge": "REJA 4: TA'LIM ISLOHOTLARI",
        "title": "4. Ta'lim Sohasidagi Tub O'zgarishlar va Inson Kapitali",
        "type": "cards",
        "cards": [
            ("Maktabgacha Ta'lim", "Maktabgacha ta'lim vazirligining tashkil etilishi.", "Qamrov darajasini 27% dan 72% dan ortiq ko'rsatkichga etkazish."),
            ("Maktablar va O'qituvchi Qadri", "O'qituvchilar ish haqini oshirish va ularni majburiy mehnatsiz qilish.", "Prezident maktablari va Ijod maktablari tarmog'ini yaratish."),
            ("Oliy Ta'lim Qamrovi", "Oliy ta'lim bilan qamrovni 9% dan 38% ga oshirish va filiallar ochish.", "OTMlarga akademik va moliyaviy mustaqillik berilishi.")
        ]
    },
    {
        "bg": "linear-gradient(135deg, #0f172a 0%, #15803d 100%)",
        "badge": "REJA 4: SOG'LIQNI SAQLASH",
        "title": "Sog'liqni Saqlash Tizimi va Aholi Salomatligi Muhofazasi",
        "type": "cards",
        "cards": [
            ("Tibbiy Xizmat Sifati", "Birinchi bo'g'in (oilaviy shifokorlik)ni rivojlantirish.", "Qishloq joylarda tibbiy maskanlar va tez yordam sifatini oshirish."),
            ("Sug'urta va Texnologiya", "Davlat tibbiy sug'urtasi tizimini bosqichma-bosqich joriy etish.", "Yuqori texnologik murakkab operatsiyalarni mahalliy darajada o'tkazish."),
            ("Aholi Salomatligi", "Ommaviy sport va sog'lom turmush tarzini targ'ib qilish.", "Aholining o'rtacha umr ko'rish davomiyligining sezilarli uzayishi.")
        ]
    },
    {
        "bg": "linear-gradient(135deg, #0f172a 0%, #1e3a8a 100%)",
        "badge": "REJA 5: TASHQI SIYOSAT",
        "title": "5. Tashqi Siyosat: Markaziy Osiyodagi Yangi Muhit",
        "type": "cards",
        "cards": [
            ("Markaziy Osiyo Ustuvorligi", "Qo'shni davlatlar bilan do'stona aloqalarni tiklash.", "Chegara va suv-energetika muammolarini tinch muloqot orqali hal etish."),
            ("Mintaqaviy Integratsiya", "Markaziy Osiyo davlatlari rahbarlarining Maslahat uchrashuvlari.", "Mintaqada o'zaro ishonch va barqarorlik muhitining qaror topishi."),
            ("Xalqaro Tashkilotlar", "BMT, ShHT, TDT va MDB doirasida faol va tashabbuskor diplomatiya.", "O'zbekistonning xalqaro sammitlar va muloqotlar markaziga aylanishi.")
        ]
    },
    {
        "bg": "linear-gradient(135deg, #0f172a 0%, #0f766e 100%)",
        "badge": "REJA 5: GLOBAL TASHABBUSLAR",
        "title": "Global Tashabbuslar, Logistika va Ekologik Diplomatiya",
        "type": "cards",
        "cards": [
            ("Transport Koridorlari", "Xitoy - Qirg'iziston - O'zbekiston temir yo'li loyihasi.", "Trans-Afg'on koridori va Janubiy Osiyoga chiqish strategiyasi."),
            ("Ekologik Diplomatiya", "Orolbo'yi mintaqasi bo'yicha global BMT rezolutsiyalari.", "'Yashil makon' umummilliy loyihasi va ekologik barqarorlik."),
            ("Muvozanatli Diplomatiya", "AQSH, Yevropa Ittifoqi, Xitoy va Rossiya bilan muvozanatli aloqalar.", "O'zbekistonning iqtisodiy diplomatiyasi va eksport geografiyasini oshirish.")
        ]
    },
    {
        "bg": "linear-gradient(135deg, #0f172a 0%, #701a75 100%)",
        "badge": "XALQARO E'TIROF",
        "title": "Yangi O'zbekiston Islohotlarining Xalqaro E'tirofi",
        "type": "cards",
        "cards": [
            ("BMT va Tashkilotlar", "O'zbekiston ilgari surgan rezolutsiyalarning BMT tomonidan qabul qilinishi.", "Inson huquqlari bo'yicha Kengashga a'zolik va faoliyat."),
            ("Nufuzli Reytinglar", "'The Economist' tomonidan O'zbekistonning 'Yil mamlakati' deb e'tirof etilishi.", "Jahon bankining 'Doing Business' va BMT indekslaridagi ko'tarilish."),
            ("Mintaqaviy Liderlik", "Markaziy Osiyoda tinchlik va hamkorlik kafolatchisi sifatida e'tirof etilishi.", "Mintaqada barqaror rivojlanish drayveriga aylanishi.")
        ]
    },
    {
        "bg": "linear-gradient(135deg, #0f172a 0%, #1e1b4b 100%)",
        "badge": "XULOSA VA ISTIQBOL",
        "title": "Xulosa: Yangi O'zbekistonning Tarixiy Roli va Kelajagi",
        "type": "cards",
        "cards": [
            ("Taraqqiyotning Ortga Qaytmasligi", "Yangi O'zbekiston strategiyasi - mamlakatimiz taraqqiyotining yangi poydevorini yaratdi.", "Islohotlar ortga qaytmas tus oldi va xalq tomonidan qo'llab-quvvatlanmoqda."),
            ("Inson Qadri - Doimiy Maydon", "Inson qadrini ulug'lash - amaliy harakatlar mezoniga aylandi.", "Har bir fuqaro uchun teng imkoniyatlar va adolatli jamiyat barpo etilmoqda."),
            ("Kelajakka Nigoh", "O'zbekiston-2030 strategiyasi doirasida yanada yuksak marralar belgilandi.", "Raqamli iqtisodiyot va yashil energetika sari sabit qadam tashlanmoqda.")
        ]
    },
    {
        "bg": "linear-gradient(135deg, #0f172a 0%, #030712 100%)",
        "badge": "E'TIBORINGIZ UCHUN RAHMAT",
        "title": "E'TIBORINGIZ UCHUN RAHMAT!",
        "subtitle": "Yangi O'zbekiston Taraqqiyot Strategiyasi va Inson Qadri Tamoyili  |  Savollar va muhokama uchun tashakkur!",
        "type": "hero"
    }
]

def generate_html(data, slide_num):
    bg = data.get("bg", "linear-gradient(135deg, #0f172a, #1e293b)")
    badge = data.get("badge", "O'ZBEKISTON TARIXI")
    title = data.get("title", "")
    subtitle = data.get("subtitle", "")
    stype = data.get("type", "cards")

    content_html = ""
    if stype == "hero":
        content_html = f"""
        <div style="margin-top: 60px;">
            <div style="font-size: 56px; font-weight: 800; line-height: 1.2; background: linear-gradient(90deg, #ffffff, #93c5fd); -webkit-background-clip: text; -webkit-text-fill-color: transparent; max-width: 1500px; margin-bottom: 30px;">{title}</div>
            <div style="font-size: 24px; color: #94a3b8; max-width: 1300px; line-height: 1.6;">{subtitle}</div>
        </div>
        """
    elif stype == "agenda":
        items_html = ""
        for num, t, d in data.get("items", []):
            items_html += f"""
            <div style="background: rgba(255,255,255,0.03); border: 1px solid rgba(255,255,255,0.08); border-radius: 20px; padding: 24px 32px; display: flex; align-items: center; gap: 24px;">
                <div style="background: linear-gradient(135deg, #38bdf8, #818cf8); color: #0f172a; font-size: 22px; font-weight: 800; width: 54px; height: 54px; border-radius: 16px; display: flex; align-items: center; justify-content: center; flex-shrink: 0;">{num}</div>
                <div>
                    <div style="font-size: 22px; font-weight: 700; color: #f8fafc; margin-bottom: 6px;">{t}</div>
                    <div style="font-size: 16px; color: #94a3b8;">{d}</div>
                </div>
            </div>
            """
        content_html = f"""
        <div>
            <div style="font-size: 42px; font-weight: 800; color: #ffffff; margin-top: 15px; margin-bottom: 30px;">{title}</div>
            <div style="display: flex; flex-direction: column; gap: 16px;">
                {items_html}
            </div>
        </div>
        """
    else:
        cards_html = ""
        for ctitle, p1, p2 in data.get("cards", []):
            cards_html += f"""
            <div style="background: rgba(255,255,255,0.03); border: 1px solid rgba(255,255,255,0.08); backdrop-filter: blur(16px); border-radius: 24px; padding: 36px; display: flex; flex-direction: column; justify-content: space-between; box-shadow: 0 20px 40px rgba(0,0,0,0.3);">
                <div>
                    <div style="display: flex; align-items: center; gap: 10px; color: #38bdf8; font-size: 14px; font-weight: 700; letter-spacing: 1.5px; text-transform: uppercase; margin-bottom: 16px;">
                        <span style="width: 8px; height: 8px; background: #38bdf8; border-radius: 50%;"></span> BO'LIM KONSEPSIYASI
                    </div>
                    <div style="font-size: 26px; font-weight: 700; color: #f8fafc; margin-bottom: 20px; line-height: 1.3;">{ctitle}</div>
                </div>
                <div style="display: flex; flex-direction: column; gap: 16px;">
                    <div style="background: rgba(255,255,255,0.04); border-left: 4px solid #38bdf8; padding: 14px 18px; border-radius: 8px; font-size: 16px; color: #cbd5e1; line-height: 1.5;">{p1}</div>
                    <div style="background: rgba(255,255,255,0.04); border-left: 4px solid #818cf8; padding: 14px 18px; border-radius: 8px; font-size: 16px; color: #cbd5e1; line-height: 1.5;">{p2}</div>
                </div>
            </div>
            """
        content_html = f"""
        <div>
            <div style="font-size: 40px; font-weight: 800; color: #ffffff; margin-top: 15px; margin-bottom: 35px;">{title}</div>
            <div style="display: grid; grid-template-columns: repeat(3, 1fr); gap: 28px; height: 500px;">
                {cards_html}
            </div>
        </div>
        """

    return f"""<!DOCTYPE html>
    <html>
    <head>
        <meta charset="UTF-8">
        <style>
            * {{ box-sizing: border-box; margin: 0; padding: 0; font-family: 'Inter', system-ui, -apple-system, sans-serif; }}
            body {{
                width: 1920px;
                height: 1080px;
                background: {bg};
                color: white;
                padding: 70px 90px;
                display: flex;
                flex-direction: column;
                justify-content: space-between;
                position: relative;
                overflow: hidden;
            }}
            .top-badge {{
                display: inline-flex;
                align-items: center;
                gap: 8px;
                background: rgba(56, 189, 248, 0.12);
                border: 1px solid rgba(56, 189, 248, 0.4);
                color: #38bdf8;
                padding: 10px 22px;
                border-radius: 30px;
                font-size: 15px;
                font-weight: 700;
                letter-spacing: 2px;
                width: fit-content;
            }}
            .footer-bar {{
                display: flex;
                align-items: center;
                justify-content: space-between;
                border-top: 1px solid rgba(255,255,255,0.1);
                padding-top: 20px;
                font-size: 16px;
                color: #64748b;
            }}
        </style>
    </head>
    <body>
        <div>
            <div class="top-badge">✦  {badge}</div>
            {content_html}
        </div>
        <div class="footer-bar">
            <div>Yangi O'zbekiston Taraqqiyot Strategiyasi va Inson Qadri Tamoyili</div>
            <div>Slayd {slide_num} / {len(SLIDES_DATA)}</div>
        </div>
    </body>
    </html>
    """

async def main():
    img_dir = r"C:\Users\user\Desktop\SlideTranslate_AI\gamma_slides_img"
    os.makedirs(img_dir, exist_ok=True)
    
    image_paths = []
    async with async_playwright() as p:
        browser = await p.chromium.launch(headless=True)
        page = await browser.new_page(viewport={'width': 1920, 'height': 1080})
        
        for idx, sdata in enumerate(SLIDES_DATA, 1):
            html = generate_html(sdata, idx)
            await page.set_content(html)
            img_p = os.path.join(img_dir, f"slide_{idx}.png")
            await page.screenshot(path=img_p)
            image_paths.append(img_p)
            print(f"Rendered Slide {idx}/{len(SLIDES_DATA)}")
            
        await browser.close()

    # Create PPTX with high-res images
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    for img_p in image_paths:
        slide = prs.slides.add_slide(blank_layout)
        slide.shapes.add_picture(img_p, 0, 0, Inches(13.333), Inches(7.5))

    out_pptx = r"C:\Users\user\Desktop\SlideTranslate_AI\generated_presentations\Yangi_Ozbekiston_Gamma_Style_UltraHD.pptx"
    os.makedirs(os.path.dirname(out_pptx), exist_ok=True)
    prs.save(out_pptx)
    print(f"GAMMA_STYLE_HD_DECK_CREATED: {out_pptx}")

if __name__ == "__main__":
    asyncio.run(main())
