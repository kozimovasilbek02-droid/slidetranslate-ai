# -*- coding: utf-8 -*-
import sys, os
sys.stdout.reconfigure(encoding='utf-8')
from pptx import Presentation

orig_path = r'C:\Users\user\Desktop\SlideTranslate_AI\temp_sessions\2abdc0f9-c72d-4b1a-a869-9cd0d9591731\original.pptx'
prs = Presentation(orig_path)
s5 = prs.slides[4]

print('Slide 5 shape count:', len(s5.shapes))
for idx, sh in enumerate(s5.shapes):
    has_tf = getattr(sh, 'has_text_frame', False)
    has_tbl = getattr(sh, 'has_table', False)
    has_ch = getattr(sh, 'has_chart', False)
    print(f'Shape {idx}: Name={sh.name}, Type={sh.shape_type}, has_tf={has_tf}, has_tbl={has_tbl}, has_ch={has_ch}')
    if has_tf and sh.text_frame:
        print('  Text:', repr(sh.text_frame.text))
    # print element xml snippet
    xml_str = sh.element.xml[:300]
    print('  XML snippet:', xml_str.replace('\n', ' '))
